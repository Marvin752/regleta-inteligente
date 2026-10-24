import hmac
import html
import io
import json
import os
import random
import re
import sqlite3
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import deque
from datetime import date, datetime, timedelta
from functools import wraps
from zoneinfo import ZoneInfo

import paho.mqtt.client as mqtt
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from flask import (Flask, Response, jsonify, redirect, render_template,
                   request, session, url_for)

load_dotenv()

MQTT_HOST = os.environ["MQTT_HOST"]
MQTT_USER = os.environ["MQTT_USER"]
MQTT_PASS = os.environ["MQTT_PASS"]
MQTT_PORT = 8883
TOPIC_ESTADO = "proyectoiot/jutiapa/estado"
TOPIC_COMANDO = "proyectoiot/jutiapa/comando"

# En Railway la base va en el volumen; en tu compu, en la carpeta del proyecto.
DB = os.path.join(os.environ.get("RAILWAY_VOLUME_MOUNT_PATH", "."), "lecturas.db")

# Los servidores en la nube usan hora UTC; las fechas se guardan en hora de Guatemala.
ZONA = ZoneInfo("America/Guatemala")

APP_USER = os.environ["APP_USER"]
APP_PASS = os.environ["APP_PASS"]

# Opcionales: si no están, la app funciona igual pero sin avisos al celular.
TELEGRAM_TOKEN = os.environ.get("TELEGRAM_TOKEN", "")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID", "")
# Asistente de IA (opcional). Groq usa el mismo formato que la API de OpenAI.
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")
# Link al panel en los avisos. Railway define RAILWAY_PUBLIC_DOMAIN al generar el dominio.
APP_URL = os.environ.get("APP_URL") or (
    "https://" + os.environ["RAILWAY_PUBLIC_DOMAIN"] if os.environ.get("RAILWAY_PUBLIC_DOMAIN") else ""
)

# El ESP32 manda una lectura cada 2 s. Si pasan más de 15 s sin nada, está desconectado.
SEGUNDOS_PARA_DESCONEXION = 15
# Para calcular energía, un hueco de más de 10 s entre lecturas no se cuenta
# (el dispositivo estaba apagado o sin conexión).
MAX_SEGUNDOS_ENTRE_LECTURAS = 10

AJUSTES_INICIALES = {"tarifa_kwh": "2.00", "voltaje": "120"}

app = Flask(__name__)
app.secret_key = os.environ["SECRET_KEY"]
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
cliente_mqtt = None

# Lo último que mandó el dispositivo, se haya guardado o no.
# De aquí sale la vista en vivo del dashboard.
recientes = deque(maxlen=20)
ultimo_mensaje = None          # cuándo llegó el último mensaje del ESP32
desconectado_avisado = False   # para no repetir el aviso de desconexión


def ahora():
    return datetime.now(ZONA).replace(tzinfo=None)


def ahora_texto():
    return ahora().isoformat(timespec="seconds")


# ---------- Base de datos ----------

def conectar():
    # timeout: varios hilos escriben (MQTT, vigilante, páginas); esperan su turno.
    return sqlite3.connect(DB, timeout=10)


def crear_tablas():
    con = conectar()
    con.executescript("""
        CREATE TABLE IF NOT EXISTS lecturas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            dispositivo TEXT NOT NULL,
            corriente REAL NOT NULL,
            limite REAL NOT NULL,
            estado TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_lecturas_fecha ON lecturas (fecha);

        CREATE TABLE IF NOT EXISTS eventos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT NOT NULL,
            tipo TEXT NOT NULL,
            detalle TEXT NOT NULL,
            usuario TEXT
        );

        CREATE TABLE IF NOT EXISTS ajustes (
            clave TEXT PRIMARY KEY,
            valor TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS horarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            hora TEXT NOT NULL,
            accion TEXT NOT NULL,
            dias TEXT NOT NULL,
            activo INTEGER NOT NULL DEFAULT 1,
            ultima_ejecucion TEXT
        );
    """)
    for clave, valor in AJUSTES_INICIALES.items():
        con.execute("INSERT OR IGNORE INTO ajustes (clave, valor) VALUES (?, ?)", (clave, valor))
    con.commit()
    con.close()


def consultar(sql, params=()):
    con = conectar()
    con.row_factory = sqlite3.Row
    filas = [dict(f) for f in con.execute(sql, params).fetchall()]
    con.close()
    return filas


def ejecutar(sql, params=()):
    con = conectar()
    cursor = con.execute(sql, params)
    con.commit()
    con.close()
    return cursor.rowcount


def guardar(lectura):
    ejecutar(
        "INSERT INTO lecturas (fecha, dispositivo, corriente, limite, estado) "
        "VALUES (?, ?, ?, ?, ?)",
        (lectura["fecha"], lectura["dispositivo"], lectura["corriente"],
         lectura["limite"], lectura["estado"]),
    )


def es_repetida(lectura, anterior):
    """Una lectura en 0 A es repetida si la anterior también estaba en 0 A
    y en el mismo estado. Esas no aportan nada al historial."""
    return (
        anterior is not None
        and lectura["corriente"] == 0
        and anterior["corriente"] == 0
        and lectura["estado"] == anterior["estado"]
    )


def leer_ajustes():
    filas = consultar("SELECT clave, valor FROM ajustes")
    ajustes = dict(AJUSTES_INICIALES)
    ajustes.update({f["clave"]: f["valor"] for f in filas})
    return {"tarifa_kwh": float(ajustes["tarifa_kwh"]), "voltaje": float(ajustes["voltaje"])}


# ---------- Bitácora y avisos ----------

def avisar_telegram(texto):
    """Manda un mensaje al celular en segundo plano. Si falla, solo se anota en consola."""
    if not (TELEGRAM_TOKEN and TELEGRAM_CHAT_ID):
        return

    def enviar():
        try:
            url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage"
            datos = urllib.parse.urlencode({
                "chat_id": TELEGRAM_CHAT_ID,
                "text": texto,
                "parse_mode": "HTML",
                "disable_web_page_preview": "true",
            }).encode()
            urllib.request.urlopen(url, data=datos, timeout=10)
        except Exception as error:
            print("No se pudo avisar por Telegram:", error)

    threading.Thread(target=enviar, daemon=True).start()


# tipo de evento -> (emoji, título, consejo). Solo estos tipos llegan al celular.
AVISOS = {
    "CORTE_LIMITE": ("⚠️", "Corte por sobrecorriente",
                     "Desconectá o revisá la carga antes de reactivar."),
    "CORTE_MANUAL": ("🔌", "Energía cortada", ""),
    "REACTIVACION": ("✅", "Energía restablecida", ""),
    "DESCONEXION": ("📴", "Dispositivo desconectado",
                    "Revisá la alimentación y el WiFi del ESP32. "
                    "La protección contra sobrecorriente sigue funcionando sin conexión."),
    "RECONEXION": ("📶", "Dispositivo en línea de nuevo", ""),
}


def texto_aviso(tipo, detalle):
    emoji, titulo, consejo = AVISOS[tipo]
    momento = ahora().strftime("%d/%m/%Y %H:%M:%S")
    lineas = [
        f"{emoji} <b>{titulo.upper()}</b>",
        "",
        html.escape(detalle),
        f"🕒 {momento}",
    ]
    if consejo:
        lineas += ["", f"<i>{html.escape(consejo)}</i>"]
    if APP_URL:
        lineas += ["", f'<a href="{html.escape(APP_URL)}">Abrir el panel</a>']
    lineas += ["", "— <b>MONITOR DE ENERGÍA</b> · PROYECTO IoT"]
    return "\n".join(lineas)


def registrar_evento(tipo, detalle, usuario=None):
    ejecutar(
        "INSERT INTO eventos (fecha, tipo, detalle, usuario) VALUES (?, ?, ?, ?)",
        (ahora_texto(), tipo, detalle, usuario),
    )
    print(f"[evento] {tipo}: {detalle}" + (f" ({usuario})" if usuario else ""))
    if tipo in AVISOS:
        avisar_telegram(texto_aviso(tipo, detalle))


# ---------- Consultas compartidas ----------

# Cada corte es la primera lectura en CORTE después de una que no lo era.
# El tipo se deduce de esa lectura: el dispositivo manda el pico que lo
# provocó (mayor al límite) si fue por sobrecorriente, y 0 A si fue manual.
SQL_CORTES = (
    "SELECT fecha, corriente, "
    "  CASE WHEN corriente > limite THEN 'LIMITE' ELSE 'MANUAL' END AS motivo "
    "FROM ("
    "  SELECT id, fecha, corriente, limite, estado,"
    "         LAG(estado) OVER (ORDER BY fecha, id) AS anterior"
    "  FROM lecturas WHERE fecha >= ? AND fecha < ?"
    ") WHERE estado = 'CORTE' AND (anterior IS NULL OR anterior <> 'CORTE') "
    "ORDER BY fecha DESC, id DESC"
)

# Energía: cada lectura vale su corriente durante el tiempo hasta la siguiente.
# amp_seg = amperios × segundos. Luego kWh = amp_seg × voltaje / 3 600 000.
SQL_ENERGIA_POR_LECTURA = f"""
    SELECT fecha,
           corriente * MAX(0, MIN(
               COALESCE((julianday(siguiente) - julianday(fecha)) * 86400, 2),
               {MAX_SEGUNDOS_ENTRE_LECTURAS}
           )) AS amp_seg
    FROM (
        SELECT fecha, corriente, LEAD(fecha) OVER (ORDER BY fecha, id) AS siguiente
        FROM lecturas WHERE fecha >= ? AND fecha < ?
    )
"""


def a_kwh(amp_seg, voltaje):
    return (amp_seg or 0) * voltaje / 3_600_000


def contar_motivos(cortes):
    limite = sum(1 for c in cortes if c["motivo"] == "LIMITE")
    return {"cortes": len(cortes), "cortes_limite": limite, "cortes_manual": len(cortes) - limite}


def motivo_corte_actual():
    """Por qué está abierto el relé ahora: 'LIMITE', 'MANUAL' o None."""
    if estado_actual() != "CORTE":
        return None
    # El corte actual empieza en la primera lectura en CORTE después de la última normal.
    # Se busca con el índice de fecha para no recorrer toda la tabla.
    ultimo = consultar(
        "SELECT corriente, limite FROM lecturas WHERE estado = 'CORTE' AND fecha >= COALESCE("
        "  (SELECT fecha FROM lecturas WHERE estado <> 'CORTE' ORDER BY fecha DESC LIMIT 1), '') "
        "ORDER BY fecha, id LIMIT 1"
    )
    if not ultimo:
        return None
    return "LIMITE" if ultimo[0]["corriente"] > ultimo[0]["limite"] else "MANUAL"


def ultima_lectura():
    """Lo más reciente que se sabe del dispositivo: primero lo que llegó
    en vivo; si la app acaba de arrancar, lo último guardado."""
    if recientes:
        return recientes[-1]
    guardada = consultar("SELECT * FROM lecturas ORDER BY fecha DESC, id DESC LIMIT 1")
    return guardada[0] if guardada else None


def estado_actual():
    """'NORMAL', 'CORTE' o None si todavía no hay datos."""
    ultima = ultima_lectura()
    return ultima["estado"] if ultima else None


def segundos_sin_datos():
    if ultimo_mensaje is not None:
        return (ahora() - ultimo_mensaje).total_seconds()
    ultima = ultima_lectura()
    if ultima is None:
        return None
    return (ahora() - datetime.fromisoformat(ultima["fecha"])).total_seconds()


# ---------- MQTT ----------

def al_conectar(client, userdata, flags, reason_code, properties):
    print("Conectado al broker", MQTT_HOST, "->", reason_code)
    client.subscribe(TOPIC_ESTADO)


ultima_orden = {"accion": None, "usuario": None, "momento": None}


def quien_ordeno(accion):
    """Usuario que mandó esta orden en los últimos 30 s, o None si vino de otro lado."""
    o = ultima_orden
    if o["accion"] == accion and o["momento"] and (ahora() - o["momento"]).total_seconds() < 30:
        return o["usuario"]
    return None


def detectar_cambios(lectura, anterior):
    """Anota en la bitácora los cortes y reactivaciones que reporta el dispositivo."""
    if anterior is None or lectura["estado"] == anterior["estado"]:
        return
    if lectura["estado"] == "CORTE":
        if lectura["corriente"] > lectura["limite"]:
            registrar_evento("CORTE_LIMITE",
                             f"La corriente llegó a {lectura['corriente']:.2f} A y el límite es "
                             f"{lectura['limite']:.2f} A. El relé se abrió para proteger el circuito.")
        else:
            quien = quien_ordeno("CORTAR")
            if quien == "programado":
                detalle = "Relé abierto por un horario programado."
            elif quien:
                detalle = f"Relé abierto desde la app por {quien}."
            else:
                detalle = "Relé abierto por orden manual."
            registrar_evento("CORTE_MANUAL", detalle)
    elif lectura["estado"] == "NORMAL" and anterior["estado"] == "CORTE":
        quien = quien_ordeno("REARMAR")
        if quien == "programado":
            detalle = "Relé reactivado por un horario programado."
        elif quien:
            detalle = f"Relé reactivado desde la app por {quien}."
        else:
            detalle = "Relé reactivado desde el dispositivo."
        registrar_evento("REACTIVACION", detalle)


def al_recibir(client, userdata, msg):
    global ultimo_mensaje, desconectado_avisado
    try:
        dato = json.loads(msg.payload.decode())
        lectura = {
            "fecha": ahora_texto(),
            "dispositivo": str(dato["dispositivo"]),
            "corriente": float(dato["corriente"]),
            "limite": float(dato["limite"]),
            "estado": str(dato["estado"]),
        }
        anterior = ultima_lectura()
        ultimo_mensaje = ahora()
        if desconectado_avisado:
            desconectado_avisado = False
            registrar_evento("RECONEXION", f"{lectura['dispositivo']} volvió a enviar datos.")
        recientes.append(lectura)          # la vista en vivo recibe todo
        detectar_cambios(lectura, anterior)
        if es_repetida(lectura, anterior):
            print("En vivo (sin guardar):", dato)
        else:
            guardar(lectura)
            print("Guardado:", dato)
    except (json.JSONDecodeError, KeyError, TypeError, ValueError) as error:
        print("Mensaje ignorado:", error)


def crear_cliente_mqtt():
    cliente = mqtt.Client(mqtt.CallbackAPIVersion.VERSION2)
    cliente.username_pw_set(MQTT_USER, MQTT_PASS)
    cliente.tls_set()
    cliente.on_connect = al_conectar
    cliente.on_message = al_recibir
    # connect_async no detiene el arranque si el broker tarda; loop_start
    # se encarga de conectar y de reconectar si se cae la conexión.
    cliente.connect_async(MQTT_HOST, MQTT_PORT)
    return cliente


def enviar_comando(accion, usuario):
    """Publica una orden para el dispositivo. El dispositivo decide si la ejecuta."""
    ultima_orden.update(accion=accion, usuario=usuario, momento=ahora())
    cliente_mqtt.publish(TOPIC_COMANDO, json.dumps({"accion": accion}), qos=1)
    texto = "Orden de corte enviada" if accion == "CORTAR" else "Orden de reactivación enviada"
    registrar_evento("ORDEN_" + accion, texto, usuario)


# ---------- Vigilante: desconexión y horarios ----------

def ejecutar_horarios():
    momento = ahora()
    hoy = momento.date().isoformat()
    hora = momento.strftime("%H:%M")
    dia = str(momento.weekday())       # 0 = lunes ... 6 = domingo
    pendientes = consultar(
        "SELECT * FROM horarios WHERE activo = 1 AND hora = ? "
        "AND (ultima_ejecucion IS NULL OR ultima_ejecucion <> ?)",
        (hora, hoy),
    )
    for h in pendientes:
        if dia not in h["dias"]:
            continue
        ejecutar("UPDATE horarios SET ultima_ejecucion = ? WHERE id = ?", (hoy, h["id"]))
        estado = estado_actual()
        if h["accion"] == "CORTAR":
            if estado == "NORMAL":
                enviar_comando("CORTAR", "programado")
            else:
                registrar_evento("HORARIO_OMITIDO", f"Corte de las {hora} omitido: el relé ya estaba abierto")
        else:
            if estado != "CORTE":
                registrar_evento("HORARIO_OMITIDO", f"Reactivación de las {hora} omitida: el relé ya estaba cerrado")
            elif motivo_corte_actual() == "LIMITE":
                # Por seguridad, un corte por sobrecorriente solo se reactiva a mano.
                registrar_evento("HORARIO_OMITIDO",
                                 f"Reactivación de las {hora} omitida: el corte fue por sobrecorriente")
            else:
                enviar_comando("REARMAR", "programado")


def vigilar_desconexion():
    global desconectado_avisado
    segundos = segundos_sin_datos()
    if ultimo_mensaje is None or desconectado_avisado or segundos is None:
        return
    if segundos > SEGUNDOS_PARA_DESCONEXION:
        desconectado_avisado = True
        registrar_evento("DESCONEXION", f"No llegan datos del ESP32 desde hace {int(segundos)} segundos.")


def vigilante():
    while True:
        try:
            vigilar_desconexion()
            ejecutar_horarios()
        except Exception as error:
            print("Error en el vigilante:", error)
        time.sleep(10)


# ---------- Sesión ----------

def requiere_login(vista):
    @wraps(vista)
    def envoltura(*args, **kwargs):
        if not session.get("usuario"):
            if request.path.startswith("/api/"):
                return jsonify({"ok": False, "mensaje": "Sesión no iniciada."}), 401
            return redirect(url_for("login"))
        return vista(*args, **kwargs)
    return envoltura


def credenciales_validas(usuario, clave):
    usuario_ok = hmac.compare_digest(usuario.encode(), APP_USER.encode())
    clave_ok = hmac.compare_digest(clave.encode(), APP_PASS.encode())
    return usuario_ok and clave_ok


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        usuario = request.form.get("usuario", "")
        clave = request.form.get("clave", "")
        if credenciales_validas(usuario, clave):
            session["usuario"] = usuario
            registrar_evento("LOGIN", "Inicio de sesión", usuario)
            return redirect(url_for("inicio"))
        registrar_evento("LOGIN_FALLIDO", f"Intento fallido con el usuario '{usuario[:40]}'")
        error = "Usuario o contraseña incorrectos."
    return render_template("login.html", error=error)


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------- Páginas ----------

def pagina(plantilla, nombre):
    return render_template(plantilla, usuario=session["usuario"], pagina=nombre)


@app.route("/")
@requiere_login
def inicio():
    return pagina("index.html", "dashboard")


@app.route("/reportes")
@requiere_login
def reportes():
    return pagina("reportes.html", "reportes")


@app.route("/bitacora")
@requiere_login
def bitacora():
    return pagina("bitacora.html", "bitacora")


@app.route("/ajustes")
@requiere_login
def ajustes():
    return pagina("ajustes.html", "ajustes")


# ---------- API del dashboard ----------

@app.route("/api/lecturas")
@requiere_login
def api_lecturas():
    # La más nueva primero. Si la app acaba de arrancar, se muestra lo guardado.
    if recientes:
        return jsonify(list(recientes)[::-1])
    return jsonify(consultar("SELECT * FROM lecturas ORDER BY fecha DESC, id DESC LIMIT 20"))


@app.route("/api/resumen")
@requiere_login
def api_resumen():
    hoy = ahora().date()
    rango = (hoy.isoformat(), (hoy + timedelta(days=1)).isoformat())
    metricas = consultar(
        "SELECT COUNT(*) AS total, MAX(corriente) AS maximo, AVG(corriente) AS promedio "
        "FROM lecturas WHERE fecha >= ? AND fecha < ?",
        rango,
    )[0]
    amp_seg = consultar(f"SELECT SUM(amp_seg) AS total FROM ({SQL_ENERGIA_POR_LECTURA})", rango)[0]["total"]
    ajustes_actuales = leer_ajustes()
    energia = a_kwh(amp_seg, ajustes_actuales["voltaje"])
    segundos = segundos_sin_datos()
    return jsonify({
        "ultima": ultima_lectura(),
        "motivo_corte": motivo_corte_actual(),
        "maximo": metricas["maximo"],
        "promedio": metricas["promedio"],
        "total": metricas["total"],
        "energia_kwh": energia,
        "costo": energia * ajustes_actuales["tarifa_kwh"],
        "en_linea": segundos is not None and segundos <= SEGUNDOS_PARA_DESCONEXION,
        "segundos_sin_datos": segundos,
        **contar_motivos(consultar(SQL_CORTES, rango)),
    })


@app.route("/api/rearmar", methods=["POST"])
@requiere_login
def api_rearmar():
    if estado_actual() != "CORTE":
        return jsonify({"ok": False, "mensaje": "El relé no está en corte."}), 409
    enviar_comando("REARMAR", session["usuario"])
    return jsonify({"ok": True, "mensaje": "Orden de reactivación enviada."})


@app.route("/api/cortar", methods=["POST"])
@requiere_login
def api_cortar():
    if estado_actual() != "NORMAL":
        return jsonify({"ok": False, "mensaje": "El relé ya está abierto o no hay datos."}), 409
    enviar_comando("CORTAR", session["usuario"])
    return jsonify({"ok": True, "mensaje": "Orden de corte enviada."})


# ---------- API de reportes ----------

PERIODOS = ("dia", "semana", "mes", "anio")


def rango_periodo(periodo, dia):
    """Devuelve (inicio, fin) del período que contiene a 'dia'. El fin no se incluye."""
    if periodo == "semana":
        inicio = dia - timedelta(days=dia.weekday())          # lunes
        fin = inicio + timedelta(days=7)
    elif periodo == "mes":
        inicio = dia.replace(day=1)
        fin = (inicio + timedelta(days=32)).replace(day=1)    # día 1 del mes siguiente
    elif periodo == "anio":
        inicio = dia.replace(month=1, day=1)
        fin = inicio.replace(year=inicio.year + 1)
    else:
        inicio = dia
        fin = dia + timedelta(days=1)
    return inicio, fin


def leer_periodo():
    periodo = request.args.get("periodo", "dia")
    if periodo not in PERIODOS:
        periodo = "dia"
    try:
        dia = date.fromisoformat(request.args.get("fecha", ""))
    except ValueError:
        dia = ahora().date()
    inicio, fin = rango_periodo(periodo, dia)
    return periodo, inicio, fin


@app.route("/api/reporte")
@requiere_login
def api_reporte():
    return jsonify(datos_reporte(*leer_periodo()))


def datos_reporte(periodo, inicio, fin):
    rango = (inicio.isoformat(), fin.isoformat())
    ajustes_actuales = leer_ajustes()
    voltaje, tarifa = ajustes_actuales["voltaje"], ajustes_actuales["tarifa_kwh"]

    resumen = consultar(
        "SELECT COUNT(*) AS total, MAX(corriente) AS maximo, "
        "AVG(corriente) AS promedio, MAX(limite) AS limite "
        "FROM lecturas WHERE fecha >= ? AND fecha < ?",
        rango,
    )[0]

    # Día -> por hora; semana y mes -> por día; año -> por mes.
    # 'grupo' sale de valores fijos, nunca de lo que manda el navegador.
    if periodo == "dia":
        grupo = "substr(fecha, 12, 2)"
    elif periodo == "anio":
        grupo = "substr(fecha, 1, 7)"
    else:
        grupo = "substr(fecha, 1, 10)"

    serie = consultar(
        f"SELECT {grupo} AS etiqueta, COUNT(*) AS lecturas, "
        "AVG(corriente) AS promedio, MAX(corriente) AS maximo "
        "FROM lecturas WHERE fecha >= ? AND fecha < ? "
        "GROUP BY etiqueta ORDER BY etiqueta",
        rango,
    )
    energia_por_grupo = {
        f["etiqueta"]: f["amp_seg"]
        for f in consultar(
            f"SELECT {grupo} AS etiqueta, SUM(amp_seg) AS amp_seg "
            f"FROM ({SQL_ENERGIA_POR_LECTURA}) GROUP BY etiqueta",
            rango,
        )
    }
    for s in serie:
        s["energia_kwh"] = a_kwh(energia_por_grupo.get(s["etiqueta"]), voltaje)
        s["costo"] = s["energia_kwh"] * tarifa

    cortes = consultar(SQL_CORTES, rango)
    resumen.update(contar_motivos(cortes))
    resumen["energia_kwh"] = sum(s["energia_kwh"] for s in serie)
    resumen["costo"] = resumen["energia_kwh"] * tarifa
    return {
        "periodo": periodo,
        "desde": inicio.isoformat(),
        "hasta": (fin - timedelta(days=1)).isoformat(),
        "tarifa_kwh": tarifa,
        "voltaje": voltaje,
        "resumen": resumen,
        "serie": serie,
        "cortes": cortes,
    }


NOMBRE_PERIODO = {"dia": "Diario", "semana": "Semanal", "mes": "Mensual", "anio": "Anual"}
AZUL = "182032"
ROJO_CLARO = "FDE2E4"
GRIS_CLARO = "F3F5F9"


def fecha_es(iso):
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}"


# Si el período tiene más lecturas que esto, la hoja "Lecturas" va resumida por minuto.
# Así el archivo sigue siendo liviano y nunca pasa del máximo de filas de Excel.
MAX_FILAS_DETALLADAS = 100_000


def celda(hoja, valor, negrita=False, color=None, fondo=None, formato=None,
          alinear=None, tamano=None, cursiva=False):
    c = WriteOnlyCell(hoja, value=valor)
    if negrita or color or tamano or cursiva:
        c.font = Font(bold=negrita, color=color, size=tamano, italic=cursiva)
    if fondo:
        c.fill = PatternFill("solid", fgColor=fondo)
    if formato:
        c.number_format = formato
    if alinear:
        c.alignment = Alignment(horizontal=alinear, vertical="center")
    return c


def hoja_tabla(libro, titulo, encabezados, filas, anchos, formatos, resaltar=None):
    """Hoja con encabezado azul, fila fija, filtros, anchos y formato de números."""
    hoja = libro.create_sheet(titulo)
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "A2"
    hoja.append([celda(hoja, e, negrita=True, color="FFFFFF", fondo=AZUL, alinear="center")
                 for e in encabezados])
    total = 0
    for fila in filas:
        fondo = ROJO_CLARO if resaltar and resaltar(fila) else None
        hoja.append([celda(hoja, v, fondo=fondo, formato=formatos[i]) for i, v in enumerate(fila)])
        total += 1
    if total:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{total + 1}"


@app.route("/api/exportar")
@requiere_login
def api_exportar():
    periodo, inicio, fin = leer_periodo()
    rep = datos_reporte(periodo, inicio, fin)
    res = rep["resumen"]
    rango_sql = (inicio.isoformat(), fin.isoformat())
    resumida = res["total"] > MAX_FILAS_DETALLADAS

    libro = Workbook(write_only=True)

    # Hoja 1: resumen
    hoja = libro.create_sheet("Resumen")
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 18
    rango = fecha_es(rep["desde"]) if rep["desde"] == rep["hasta"] else \
        f"Del {fecha_es(rep['desde'])} al {fecha_es(rep['hasta'])}"
    hoja.append([celda(hoja, "Monitor de energía · Proyecto IoT", negrita=True, tamano=16, color=AZUL)])
    hoja.append([celda(hoja, f"Reporte {NOMBRE_PERIODO[periodo].lower()} · {rango}",
                       cursiva=True, color="6E788C")])
    hoja.append([])
    datos = [
        ("Lecturas registradas", res["total"], "#,##0"),
        ("Corriente máxima (A)", res["maximo"], "0.00"),
        ("Corriente promedio (A)", res["promedio"], "0.00"),
        ("Energía estimada (kWh)", res["energia_kwh"], "0.000"),
        ("Costo estimado (Q)", res["costo"], '"Q"#,##0.00'),
        ("Cortes totales", res["cortes"], "0"),
        ("   Por sobrecorriente", res["cortes_limite"], "0"),
        ("   Manuales o programados", res["cortes_manual"], "0"),
        ("Tarifa usada (Q por kWh)", rep["tarifa_kwh"], '"Q"0.00'),
        ("Voltaje usado (V)", rep["voltaje"], "0"),
    ]
    for i, (nombre, valor, formato) in enumerate(datos):
        fondo = GRIS_CLARO if i % 2 == 0 else None
        hoja.append([
            celda(hoja, nombre, negrita=not nombre.startswith(" "), fondo=fondo),
            celda(hoja, valor, fondo=fondo, formato=formato, alinear="right"),
        ])
    hoja.append([])
    hoja.append([celda(hoja, "La energía es una estimación: el sensor mide corriente y se "
                       "multiplica por el voltaje indicado.", cursiva=True, tamano=9, color="6E788C")])
    if resumida:
        hoja.append([celda(hoja, f"El período tiene {res['total']:,} lecturas: la hoja Lecturas "
                           "muestra el promedio por minuto.", cursiva=True, tamano=9, color="6E788C")])
    hoja.append([celda(hoja, f"Generado el {ahora().strftime('%d/%m/%Y %H:%M')} por {session['usuario']}",
                       tamano=9, color="6E788C")])

    # Hoja 2: detalle por hora, día o mes
    grupo = {"dia": "Hora", "anio": "Mes"}.get(periodo, "Día")

    def etiqueta(e):
        if periodo == "dia":
            return f"{e}:00"
        return e if periodo == "anio" else fecha_es(e)

    hoja_tabla(
        libro, "Detalle",
        [grupo, "Lecturas", "Promedio (A)", "Máximo (A)", "Energía (kWh)", "Costo (Q)"],
        ([etiqueta(x["etiqueta"]), x["lecturas"], x["promedio"], x["maximo"], x["energia_kwh"], x["costo"]]
         for x in rep["serie"]),
        [14, 12, 14, 14, 15, 13],
        [None, "#,##0", "0.00", "0.00", "0.000", '"Q"#,##0.00'],
    )

    # Hoja 3: cortes
    hoja_tabla(
        libro, "Cortes",
        ["Fecha", "Hora", "Tipo", "Corriente (A)"],
        ([fecha_es(c["fecha"]), c["fecha"][11:],
          "Sobrecorriente" if c["motivo"] == "LIMITE" else "Manual o programado", c["corriente"]]
         for c in reversed(rep["cortes"])),
        [13, 11, 22, 15],
        [None, None, None, "0.00"],
        resaltar=lambda f: f[2] == "Sobrecorriente",
    )

    # Hoja 4: lecturas (todas, o promedio por minuto si son demasiadas)
    con = conectar()
    if resumida:
        cursor = con.execute(
            "SELECT substr(fecha, 1, 16) AS minuto, COUNT(*), AVG(corriente), MAX(corriente), "
            "MAX(estado = 'CORTE') FROM lecturas WHERE fecha >= ? AND fecha < ? "
            "GROUP BY minuto ORDER BY minuto", rango_sql)
        hoja_tabla(
            libro, "Lecturas",
            ["Fecha", "Hora", "Lecturas", "Promedio (A)", "Máximo (A)", "Hubo corte"],
            ([fecha_es(m), m[11:], n, prom, maxi, "Sí" if corte else "No"]
             for m, n, prom, maxi, corte in cursor),
            [13, 9, 11, 14, 13, 12],
            [None, None, "0", "0.00", "0.00", None],
            resaltar=lambda f: f[5] == "Sí",
        )
    else:
        cursor = con.execute(
            "SELECT fecha, dispositivo, corriente, limite, estado FROM lecturas "
            "WHERE fecha >= ? AND fecha < ? ORDER BY fecha, id", rango_sql)
        hoja_tabla(
            libro, "Lecturas",
            ["Fecha", "Hora", "Dispositivo", "Corriente (A)", "Límite (A)", "Estado"],
            ([fecha_es(f), f[11:], d, c, l, e] for f, d, c, l, e in cursor),
            [13, 11, 20, 15, 12, 11],
            [None, None, None, "0.00", "0.00", None],
            resaltar=lambda f: f[5] == "CORTE",
        )
    con.close()

    salida = io.BytesIO()
    libro.save(salida)
    nombre = f"reporte-energia-{periodo}-{inicio.isoformat()}.xlsx"
    return Response(
        salida.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@app.route("/api/mapa")
@requiere_login
def api_mapa():
    return jsonify(matriz_consumo())


def matriz_consumo():
    """Energía (Wh) por día de la semana y hora, de los últimos 30 días."""
    hoy = ahora().date()
    rango = ((hoy - timedelta(days=29)).isoformat(), (hoy + timedelta(days=1)).isoformat())
    voltaje = leer_ajustes()["voltaje"]
    celdas = consultar(
        "SELECT CAST(strftime('%w', fecha) AS INTEGER) AS dia_sql, "
        "CAST(substr(fecha, 12, 2) AS INTEGER) AS hora, SUM(amp_seg) AS amp_seg "
        f"FROM ({SQL_ENERGIA_POR_LECTURA}) GROUP BY dia_sql, hora",
        rango,
    )
    matriz = [[0.0] * 24 for _ in range(7)]          # fila 0 = lunes
    for c in celdas:
        dia = (c["dia_sql"] + 6) % 7                 # SQLite: 0 = domingo
        matriz[dia][c["hora"]] = a_kwh(c["amp_seg"], voltaje) * 1000
    return {"desde": rango[0], "hasta": hoy.isoformat(), "wh": matriz}


# ---------- API de bitácora ----------

@app.route("/api/eventos")
@requiere_login
def api_eventos():
    return jsonify(consultar("SELECT * FROM eventos ORDER BY fecha DESC, id DESC LIMIT 1000"))


# ---------- API de ajustes y horarios ----------

def requiere_json():
    if not request.is_json:
        return jsonify({"ok": False, "mensaje": "Formato inválido."}), 400
    return None


@app.route("/api/ajustes", methods=["GET", "POST"])
@requiere_login
def api_ajustes():
    if request.method == "GET":
        return jsonify({**leer_ajustes(), "telegram": bool(TELEGRAM_TOKEN and TELEGRAM_CHAT_ID)})
    if (error := requiere_json()):
        return error
    datos = request.get_json()
    try:
        tarifa = float(datos.get("tarifa_kwh"))
        voltaje = float(datos.get("voltaje"))
    except (TypeError, ValueError):
        return jsonify({"ok": False, "mensaje": "Los valores deben ser números."}), 400
    if not (0 <= tarifa <= 100 and 90 <= voltaje <= 260):
        return jsonify({"ok": False, "mensaje": "Tarifa entre 0 y 100, voltaje entre 90 y 260."}), 400
    for clave, valor in (("tarifa_kwh", tarifa), ("voltaje", voltaje)):
        ejecutar("INSERT OR REPLACE INTO ajustes (clave, valor) VALUES (?, ?)", (clave, str(valor)))
    registrar_evento("AJUSTES", f"Tarifa Q{tarifa:.2f}/kWh, voltaje {voltaje:.0f} V", session["usuario"])
    return jsonify({"ok": True, "mensaje": "Ajustes guardados."})


@app.route("/api/horarios", methods=["GET", "POST"])
@requiere_login
def api_horarios():
    if request.method == "GET":
        return jsonify(consultar("SELECT * FROM horarios ORDER BY hora"))
    if (error := requiere_json()):
        return error
    datos = request.get_json()
    hora = str(datos.get("hora", ""))
    accion = str(datos.get("accion", ""))
    dias = "".join(sorted(set(str(datos.get("dias", "")))))
    if not re.fullmatch(r"([01]\d|2[0-3]):[0-5]\d", hora):
        return jsonify({"ok": False, "mensaje": "Hora inválida."}), 400
    if accion not in ("CORTAR", "REARMAR"):
        return jsonify({"ok": False, "mensaje": "Acción inválida."}), 400
    if not dias or not re.fullmatch(r"[0-6]+", dias):
        return jsonify({"ok": False, "mensaje": "Elegí al menos un día."}), 400
    ejecutar("INSERT INTO horarios (hora, accion, dias) VALUES (?, ?, ?)", (hora, accion, dias))
    nombre = "Corte programado" if accion == "CORTAR" else "Reactivación programada"
    registrar_evento("HORARIO", f"{nombre} a las {hora}", session["usuario"])
    return jsonify({"ok": True, "mensaje": "Horario agregado."})


@app.route("/api/horarios/<int:id_horario>", methods=["PATCH", "DELETE"])
@requiere_login
def api_horario(id_horario):
    if request.method == "DELETE":
        cambiados = ejecutar("DELETE FROM horarios WHERE id = ?", (id_horario,))
        mensaje = "Horario eliminado."
    else:
        if (error := requiere_json()):
            return error
        activo = 1 if request.get_json().get("activo") else 0
        cambiados = ejecutar("UPDATE horarios SET activo = ? WHERE id = ?", (activo, id_horario))
        mensaje = "Horario activado." if activo else "Horario pausado."
    if not cambiados:
        return jsonify({"ok": False, "mensaje": "Ese horario no existe."}), 404
    return jsonify({"ok": True, "mensaje": mensaje})


@app.route("/api/telegram/prueba", methods=["POST"])
@requiere_login
def api_telegram_prueba():
    if not (TELEGRAM_TOKEN and TELEGRAM_CHAT_ID):
        return jsonify({"ok": False, "mensaje": "Telegram no está configurado."}), 400
    avisar_telegram(
        "🔔 <b>MENSAJE DE PRUEBA</b>\n\n"
        "Si ves esto, los avisos están funcionando. Te voy a escribir cuando:\n"
        "⚠️ El relé corte por sobrecorriente\n"
        "🔌 Alguien corte la energía desde la app\n"
        "✅ Se restablezca la energía\n"
        "📴 El ESP32 se desconecte\n\n"
        "— <b>MONITOR DE ENERGÍA</b> · PROYECTO IoT"
    )
    return jsonify({"ok": True, "mensaje": "Mensaje de prueba enviado."})


# ---------- Asistente de IA ----------
# La IA no lee la base directamente: llama a estas funciones, que usan las mismas
# consultas que los reportes. Solo pueden leer; no hay forma de mover el relé desde aquí.

DIAS_SEMANA = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def _redondear(valor, decimales=3):
    return round(valor, decimales) if isinstance(valor, float) else valor


def herramienta_resumen(periodo="dia", fecha=None):
    try:
        dia = date.fromisoformat(fecha) if fecha else ahora().date()
    except ValueError:
        dia = ahora().date()
    periodo = periodo if periodo in PERIODOS else "dia"
    rep = datos_reporte(periodo, *rango_periodo(periodo, dia))
    return {
        "periodo": periodo,
        "desde": rep["desde"],
        "hasta": rep["hasta"],
        "resumen": {k: _redondear(v) for k, v in rep["resumen"].items()},
        "detalle": [{k: _redondear(v) for k, v in fila.items()} for fila in rep["serie"]],
        "unidades": "corriente en A, energia_kwh en kWh, costo en quetzales",
    }


def herramienta_cortes(periodo="mes", fecha=None, tipo=None):
    try:
        dia = date.fromisoformat(fecha) if fecha else ahora().date()
    except ValueError:
        dia = ahora().date()
    periodo = periodo if periodo in PERIODOS else "mes"
    inicio, fin = rango_periodo(periodo, dia)
    cortes = consultar(SQL_CORTES, (inicio.isoformat(), fin.isoformat()))
    if tipo in ("LIMITE", "MANUAL"):
        cortes = [c for c in cortes if c["motivo"] == tipo]
    por_hora, por_dia = {}, {}
    for c in cortes:
        momento = datetime.fromisoformat(c["fecha"])
        por_hora[momento.hour] = por_hora.get(momento.hour, 0) + 1
        nombre = DIAS_SEMANA[momento.weekday()]
        por_dia[nombre] = por_dia.get(nombre, 0) + 1
    return {
        "desde": inicio.isoformat(),
        "hasta": (fin - timedelta(days=1)).isoformat(),
        "total": len(cortes),
        **contar_motivos(cortes),
        "por_hora_del_dia": dict(sorted(por_hora.items())),
        "por_dia_de_la_semana": por_dia,
        "lista": cortes[:80],
        "nota": "LIMITE = sobrecorriente (corriente = pico que lo provocó); MANUAL = desde la app o programado",
    }


def herramienta_patron():
    datos = matriz_consumo()
    celdas = [(DIAS_SEMANA[d], h, wh) for d, horas in enumerate(datos["wh"]) for h, wh in enumerate(horas)]
    por_dia = {DIAS_SEMANA[d]: round(sum(horas), 1) for d, horas in enumerate(datos["wh"])}
    por_hora = {h: round(sum(datos["wh"][d][h] for d in range(7)), 1) for h in range(24)}
    return {
        "desde": datos["desde"],
        "hasta": datos["hasta"],
        "wh_por_dia_de_la_semana": por_dia,
        "wh_por_hora_del_dia": por_hora,
        "horas_de_mayor_consumo": [
            {"dia": d, "hora": h, "wh": round(wh, 1)}
            for d, h, wh in sorted(celdas, key=lambda c: c[2], reverse=True)[:8]
        ],
    }


def herramienta_estado():
    segundos = segundos_sin_datos()
    ultima = ultima_lectura()
    hoy = herramienta_resumen("dia")["resumen"]
    return {
        "ultima_lectura": ultima,
        "dispositivo_en_linea": segundos is not None and segundos <= SEGUNDOS_PARA_DESCONEXION,
        "segundos_sin_datos": None if segundos is None else int(segundos),
        "motivo_corte_actual": motivo_corte_actual(),
        "hoy": hoy,
        "ajustes": leer_ajustes(),
    }


HERRAMIENTAS = {
    "resumen_periodo": herramienta_resumen,
    "listar_cortes": herramienta_cortes,
    "patron_consumo": herramienta_patron,
    "estado_actual": herramienta_estado,
}

PARAM_PERIODO = {
    "periodo": {"type": "string", "enum": ["dia", "semana", "mes", "anio"],
                "description": "Período a consultar"},
    "fecha": {"type": "string",
              "description": "Una fecha YYYY-MM-DD dentro del período. Si se omite, se usa hoy."},
}

DEFINICION_HERRAMIENTAS = [
    {"type": "function", "function": {
        "name": "resumen_periodo",
        "description": "Lecturas, corriente máxima y promedio, energía (kWh), costo (Q) y cortes de un "
                       "día, semana, mes o año, con el detalle por hora (día), por día (semana o mes) "
                       "o por mes (año).",
        "parameters": {"type": "object", "properties": PARAM_PERIODO, "required": ["periodo"]}}},
    {"type": "function", "function": {
        "name": "listar_cortes",
        "description": "Cortes del relé en un período: cuántos, de qué tipo, a qué horas y qué días "
                       "de la semana, y la lista con fecha y corriente.",
        "parameters": {"type": "object", "properties": {
            **PARAM_PERIODO,
            "tipo": {"type": "string", "enum": ["LIMITE", "MANUAL"],
                     "description": "Filtrar por tipo de corte (opcional)"}},
            "required": ["periodo"]}}},
    {"type": "function", "function": {
        "name": "patron_consumo",
        "description": "Energía de los últimos 30 días agrupada por día de la semana y por hora "
                       "del día, y las horas de mayor consumo.",
        "parameters": {"type": "object", "properties": {}}}},
    {"type": "function", "function": {
        "name": "estado_actual",
        "description": "Estado en este momento: última lectura, si el ESP32 está en línea, "
                       "si el relé está cortado y por qué, el consumo de hoy y la tarifa.",
        "parameters": {"type": "object", "properties": {}}}},
]


def instrucciones_sistema():
    momento = ahora()
    ajustes_actuales = leer_ajustes()
    return (
        "Te llamás Rem y sos la asistente del Monitor de energía, un proyecto IoT universitario en Jutiapa, "
        "Guatemala. También sos la que avisa por Telegram cuando hay cortes o desconexiones. "
        "Un ESP32 con un sensor ACS712 mide la corriente de una regleta cada 2 segundos y un relé "
        f"corta la energía si pasa de 2.5 A. Hoy es {DIAS_SEMANA[momento.weekday()]} "
        f"{momento.strftime('%Y-%m-%d')} y son las {momento.strftime('%H:%M')} (hora de Guatemala). "
        f"La tarifa configurada es Q{ajustes_actuales['tarifa_kwh']:.2f} por kWh y el voltaje "
        f"{ajustes_actuales['voltaje']:.0f} V; la energía es una estimación a partir de la corriente.\n"
        "Reglas:\n"
        "- Para cualquier dato, usá las herramientas. Nunca inventes cifras.\n"
        "- Respondé en español, breve y claro, con las cifras y sus unidades (A, kWh, Q).\n"
        "- Si comparás períodos, consultá cada uno.\n"
        "- Podés dar recomendaciones de ahorro o seguridad basadas en los datos.\n"
        "- No podés cortar ni reactivar el relé: eso se hace con los botones del Dashboard.\n"
        "- Si no hay datos para lo que te preguntan, decilo."
    )


def llamar_groq(mensajes):
    cuerpo = json.dumps({
        "model": GROQ_MODEL,
        "messages": mensajes,
        "tools": DEFINICION_HERRAMIENTAS,
        "tool_choice": "auto",
        "temperature": 0.2,
        "max_tokens": 900,
    }).encode()
    pedido = urllib.request.Request(
        "https://api.groq.com/openai/v1/chat/completions",
        data=cuerpo,
        headers={"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json",
                 "User-Agent": "monitor-energia/1.0"},
    )
    with urllib.request.urlopen(pedido, timeout=60) as respuesta:
        return json.loads(respuesta.read())["choices"][0]["message"]


@app.route("/api/chat", methods=["POST"])
@requiere_login
def api_chat():
    if not GROQ_API_KEY:
        return jsonify({"ok": False, "mensaje": "Rem no está configurada (falta GROQ_API_KEY)."}), 400
    historial = (request.get_json(silent=True) or {}).get("mensajes", [])
    mensajes = [{"role": "system", "content": instrucciones_sistema()}]
    for m in historial[-10:]:
        if isinstance(m, dict) and m.get("role") in ("user", "assistant") and isinstance(m.get("content"), str):
            mensajes.append({"role": m["role"], "content": m["content"][:2000]})
    if len(mensajes) == 1 or mensajes[-1]["role"] != "user":
        return jsonify({"ok": False, "mensaje": "Escribí una pregunta."}), 400

    try:
        for _ in range(6):
            respuesta = llamar_groq(mensajes)
            llamadas = respuesta.get("tool_calls") or []
            if not llamadas:
                return jsonify({"ok": True, "respuesta": respuesta.get("content") or "No tengo una respuesta."})
            mensajes.append({"role": "assistant", "content": respuesta.get("content") or "",
                             "tool_calls": llamadas})
            for llamada in llamadas:
                nombre = llamada["function"]["name"]
                try:
                    argumentos = json.loads(llamada["function"].get("arguments") or "{}")
                    resultado = HERRAMIENTAS[nombre](**argumentos)
                except Exception as error:
                    resultado = {"error": f"No se pudo consultar: {error}"}
                mensajes.append({"role": "tool", "tool_call_id": llamada["id"],
                                 "content": json.dumps(resultado, ensure_ascii=False, default=str)[:20000]})
        return jsonify({"ok": True, "respuesta": "La pregunta necesitó demasiadas consultas. Probá algo más concreto."})
    except urllib.error.HTTPError as error:
        detalle = error.read().decode(errors="ignore")[:300]
        print("Error de Groq:", error.code, detalle)
        if error.code == 429:
            mensaje = "Se alcanzó el límite de preguntas por minuto. Esperá un momento y volvé a intentar."
        elif error.code == 401:
            mensaje = "La clave de Groq no es válida."
        else:
            mensaje = "Rem no pudo responder. Intentá de nuevo."
        return jsonify({"ok": False, "mensaje": mensaje}), 502
    except Exception as error:
        print("Error del asistente:", error)
        return jsonify({"ok": False, "mensaje": "Rem no pudo responder. Intentá de nuevo."}), 502


# ---------- Datos de demostración ----------
# Genera un historial parecido al consumo real medido (0.3 a 0.5 A de base, picos
# de hasta ~1 A, cortes manuales ocasionales) para probar la app con meses de datos.
# Los rangos de id insertados se guardan en "ajustes" para poder borrarlos después.

DISPOSITIVO_DEMO = "esp32-jutiapa-01"
LIMITE_DEMO = 2.5
# Corriente media por hora del día, sacada del reporte real del 23/10
PERFIL_HORARIO = {7: .30, 8: .32, 9: .34, 10: .36, 11: .38, 12: .40, 13: .42, 14: .45,
                  15: .41, 16: .40, 17: .48, 18: .36, 19: .40, 20: .39, 21: .34}
progreso_demo = {"activo": False, "dia": 0, "total": 0, "lecturas": 0, "mensaje": ""}


def lecturas_de_un_dia(dia, azar):
    """Devuelve (lecturas, eventos) simulados para un día."""
    filas, eventos = [], []
    if azar.random() < 0.06:                       # algunos días no se encendió
        return filas, eventos
    fin_de_semana = dia.weekday() >= 5
    inicio = datetime.combine(dia, datetime.min.time()) + timedelta(
        hours=9 if fin_de_semana else 7, minutes=azar.randint(0, 50))
    fin = datetime.combine(dia, datetime.min.time()) + timedelta(
        hours=20 if fin_de_semana else 21, minutes=azar.randint(0, 59))
    factor_dia = azar.uniform(0.85, 1.15)

    # Cortes del día: manuales casi a diario, por sobrecorriente de vez en cuando
    cortes = []
    for _ in range(azar.choice([0, 1, 1, 1, 2, 2, 3])):
        cortes.append((inicio + (fin - inicio) * azar.random(), "MANUAL",
                       timedelta(seconds=azar.randint(30, 1200))))
    if azar.random() < 0.12:
        cortes.append((inicio + (fin - inicio) * azar.random(), "LIMITE",
                       timedelta(seconds=azar.randint(120, 900))))
    cortes.sort()

    momento, pico_hasta, pico_extra = inicio, None, 0.0
    i_corte = 0
    while momento < fin:
        if i_corte < len(cortes) and momento >= cortes[i_corte][0]:
            _, motivo, duracion = cortes[i_corte]
            i_corte += 1
            if motivo == "LIMITE":
                pico = round(azar.uniform(2.55, 2.95), 2)
                filas.append((momento.isoformat(timespec="seconds"), pico, "CORTE"))
                eventos.append((momento, "CORTE_LIMITE",
                                f"La corriente llegó a {pico:.2f} A y el límite es {LIMITE_DEMO:.2f} A. "
                                "El relé se abrió para proteger el circuito.", None))
                momento += timedelta(seconds=2)
            else:
                eventos.append((momento - timedelta(seconds=1), "ORDEN_CORTAR",
                                "Orden de corte enviada", "admin"))
                eventos.append((momento, "CORTE_MANUAL", "Relé abierto desde la app por admin.", None))
            # Mientras dura el corte solo se guarda el primer cero (igual que la app real)
            filas.append((momento.isoformat(timespec="seconds"), 0.0, "CORTE"))
            momento += duracion
            eventos.append((momento - timedelta(seconds=1), "ORDEN_REARMAR",
                            "Orden de reactivación enviada", "admin"))
            eventos.append((momento, "REACTIVACION", "Relé reactivado desde la app por admin.", None))
            continue

        # Aparatos que se encienden un rato y suben el consumo
        if pico_hasta is None and azar.random() < 0.0004:
            pico_hasta = momento + timedelta(seconds=azar.randint(120, 600))
            pico_extra = azar.uniform(0.2, 0.6)
        if pico_hasta is not None and momento >= pico_hasta:
            pico_hasta, pico_extra = None, 0.0

        base = PERFIL_HORARIO.get(momento.hour, 0.32) * factor_dia
        corriente = max(0.15, base + pico_extra + azar.gauss(0, 0.05))
        filas.append((momento.isoformat(timespec="seconds"), round(corriente, 2), "NORMAL"))
        momento += timedelta(seconds=2 if azar.random() < 0.8 else 3)
    return filas, eventos


def generar_demo(dias):
    hoy = ahora().date()
    azar = random.Random(dias)
    progreso_demo.update(activo=True, dia=0, total=dias, lecturas=0, mensaje="Generando…")
    try:
        con = conectar()
        primer_id = con.execute("SELECT COALESCE(MAX(id), 0) + 1 FROM lecturas").fetchone()[0]
        primer_evento = con.execute("SELECT COALESCE(MAX(id), 0) + 1 FROM eventos").fetchone()[0]
        for n in range(dias, 0, -1):
            filas, eventos = lecturas_de_un_dia(hoy - timedelta(days=n), azar)
            con.executemany(
                "INSERT INTO lecturas (fecha, dispositivo, corriente, limite, estado) VALUES (?, ?, ?, ?, ?)",
                [(f, DISPOSITIVO_DEMO, c, LIMITE_DEMO, e) for f, c, e in filas],
            )
            con.executemany(
                "INSERT INTO eventos (fecha, tipo, detalle, usuario) VALUES (?, ?, ?, ?)",
                [(m.isoformat(timespec="seconds"), t, d, u) for m, t, d, u in eventos],
            )
            con.commit()
            progreso_demo["dia"] += 1
            progreso_demo["lecturas"] += len(filas)
        ultimo_id = con.execute("SELECT COALESCE(MAX(id), 0) FROM lecturas").fetchone()[0]
        ultimo_evento = con.execute("SELECT COALESCE(MAX(id), 0) FROM eventos").fetchone()[0]
        rangos = json.loads(leer_ajuste("demo_rangos", "[]"))
        rangos.append([primer_id, ultimo_id, primer_evento, ultimo_evento, hoy.isoformat()])
        con.execute("INSERT OR REPLACE INTO ajustes (clave, valor) VALUES ('demo_rangos', ?)",
                    (json.dumps(rangos),))
        con.commit()
        con.close()
        progreso_demo["mensaje"] = f"Listo: {progreso_demo['lecturas']:,} lecturas generadas."
    except Exception as error:
        progreso_demo["mensaje"] = f"Error: {error}"
        print("Error al generar datos de demostración:", error)
    finally:
        progreso_demo["activo"] = False


def leer_ajuste(clave, por_defecto):
    filas = consultar("SELECT valor FROM ajustes WHERE clave = ?", (clave,))
    return filas[0]["valor"] if filas else por_defecto


@app.route("/api/datos")
@requiere_login
def api_datos():
    total = consultar("SELECT COUNT(*) AS n, MIN(fecha) AS desde FROM lecturas")[0]
    rangos = json.loads(leer_ajuste("demo_rangos", "[]"))
    demo = sum(
        consultar("SELECT COUNT(*) AS n FROM lecturas WHERE id BETWEEN ? AND ? AND fecha < ?",
                  (a, b, hasta))[0]["n"]
        for a, b, _, _, hasta in rangos
    )
    tamano = sum(os.path.getsize(DB + extra) for extra in ("", "-wal") if os.path.exists(DB + extra))
    return jsonify({
        "lecturas": total["n"],
        "desde": total["desde"],
        "lecturas_demo": demo,
        "tamano_mb": tamano / 1_048_576,
        "progreso": progreso_demo,
    })


@app.route("/api/demo/generar", methods=["POST"])
@requiere_login
def api_demo_generar():
    if progreso_demo["activo"]:
        return jsonify({"ok": False, "mensaje": "Ya se están generando datos."}), 409
    if json.loads(leer_ajuste("demo_rangos", "[]")):
        return jsonify({"ok": False, "mensaje": "Ya hay datos de demostración. Borralos primero."}), 409
    try:
        dias = int((request.get_json(silent=True) or {}).get("dias", 60))
    except (TypeError, ValueError):
        dias = 60
    dias = max(7, min(dias, 120))
    threading.Thread(target=generar_demo, args=(dias,), daemon=True).start()
    registrar_evento("DEMO", f"Generando {dias} días de datos de demostración", session["usuario"])
    return jsonify({"ok": True, "mensaje": f"Generando {dias} días de datos…"})


@app.route("/api/demo/borrar", methods=["POST"])
@requiere_login
def api_demo_borrar():
    if progreso_demo["activo"]:
        return jsonify({"ok": False, "mensaje": "Esperá a que termine de generar."}), 409
    rangos = json.loads(leer_ajuste("demo_rangos", "[]"))
    if not rangos:
        return jsonify({"ok": False, "mensaje": "No hay datos de demostración."}), 404
    con = conectar()
    borradas = 0
    # Solo se borra lo anterior al día en que se generó: si el ESP32 mandó datos
    # mientras tanto, esas lecturas reales son de ese día y se conservan.
    for a, b, ea, eb, hasta in rangos:
        borradas += con.execute("DELETE FROM lecturas WHERE id BETWEEN ? AND ? AND fecha < ?",
                                (a, b, hasta)).rowcount
        con.execute("DELETE FROM eventos WHERE id BETWEEN ? AND ? AND fecha < ?", (ea, eb, hasta))
    con.execute("DELETE FROM ajustes WHERE clave = 'demo_rangos'")
    con.commit()
    con.execute("VACUUM")            # devuelve el espacio al disco
    con.close()
    registrar_evento("DEMO", f"Datos de demostración borrados ({borradas:,} lecturas)", session["usuario"])
    return jsonify({"ok": True, "mensaje": f"Se borraron {borradas:,} lecturas de demostración."})


# ---------- Arranque ----------

# Esto corre al cargar el archivo, tanto con "python app.py" como con gunicorn.
crear_tablas()
cliente_mqtt = crear_cliente_mqtt()
cliente_mqtt.loop_start()
threading.Thread(target=vigilante, daemon=True).start()

if __name__ == "__main__":
    app.run(port=int(os.environ.get("PORT", 5000)))